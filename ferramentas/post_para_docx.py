"""Transforma o texto de um post do LinkedIn em um arquivo .docx e confere as regras do post.

Uso:
    python ferramentas/post_para_docx.py linkedin/AAAA-MM-DD.txt linkedin/AAAA-MM-DD.docx

O arquivo .txt tem o post e, depois de uma linha escrita exatamente "Fontes:", os links usados.
Só o post conta para o limite de caracteres. Se alguma regra for quebrada, o script grava o
.docx mesmo assim, lista os avisos e termina com código 1, para o agente corrigir e rodar de novo.
"""

import re
import sys

from docx import Document
from docx.shared import Pt

LIMITE = 1300
TRAVESSOES = ("\u2014", "\u2013")  # travessão e meia-risca
EMOJI = re.compile(
    "[\U0001F300-\U0001FAFF\U0001F900-\U0001F9FF\U0001F1E6-\U0001F1FF\u2600-\u27BF\u2B00-\u2BFF]"
)
LISTA = re.compile(r"^\s*(?:[-*\u2022]|\d+[.)])\s", re.M)


def separar(texto):
    partes = re.split(r"^Fontes:\s*$", texto, maxsplit=1, flags=re.M)
    post = partes[0].strip()
    fontes = partes[1].strip() if len(partes) > 1 else ""
    return post, fontes


def conferir(post, fontes):
    avisos = []
    n = len(post)
    if n > LIMITE:
        avisos.append(f"o post tem {n} caracteres; o limite é {LIMITE}")
    if any(t in post for t in TRAVESSOES):
        avisos.append("o post tem travessão ou meia-risca")
    if EMOJI.search(post):
        avisos.append("o post tem emoji")
    if LISTA.search(post):
        avisos.append("o post tem lista (linha começando com hífen, asterisco ou número)")
    if "\n\n\n" in post:
        avisos.append("o post tem mais de uma linha em branco seguida")
    if "#" in post:
        avisos.append("o post tem cerquilha (hashtag ou título de Markdown)")
    if "**" in post or "__" in post:
        avisos.append("o post tem marcação de negrito de Markdown")
    if not fontes:
        avisos.append("faltou a seção 'Fontes:' com os links das notícias usadas")
    elif "http" not in fontes:
        avisos.append("a seção 'Fontes:' não tem nenhum link")
    return n, avisos


def gravar(post, fontes, saida):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    for paragrafo in re.split(r"\n\s*\n", post):
        doc.add_paragraph(paragrafo.strip())
    if fontes:
        doc.add_paragraph("")
        doc.add_paragraph("Fontes (para o primeiro comentário):")
        for linha in fontes.splitlines():
            if linha.strip():
                doc.add_paragraph(linha.strip())
    doc.save(saida)


def main():
    # No Windows, o terminal nem sempre está em UTF-8; isto evita acento quebrado nos avisos.
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except AttributeError:
        pass
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    entrada, saida = sys.argv[1], sys.argv[2]
    with open(entrada, encoding="utf-8") as f:
        texto = f.read().replace("\r\n", "\n").strip()
    post, fontes = separar(texto)
    n, avisos = conferir(post, fontes)
    gravar(post, fontes, saida)
    print(f"Gravado: {saida}")
    print(f"Caracteres do post, com espaços: {n} / {LIMITE}")
    print(f"Parágrafos: {len(re.split(r'\n\s*\n', post))}")
    if avisos:
        print("AVISOS, corrija e rode de novo:")
        for aviso in avisos:
            print(f" - {aviso}")
        sys.exit(1)
    print("Sem avisos.")


if __name__ == "__main__":
    main()
